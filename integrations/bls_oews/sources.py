"""Retrieve OEWS median hourly wages from the BLS Public Data API or an official file.

Nothing here invents or substitutes data: an occupation without a valid value for
the Atlanta area is returned with ``median_hourly_wage=None`` and a reason.
"""

from __future__ import annotations

import csv
import re
import time
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable

import requests

from . import config


class SourceError(RuntimeError):
    """The wage source could not be retrieved or understood."""


@dataclass(frozen=True)
class WageObservation:
    soc_code: str
    area_code: str
    year: int
    median_hourly_wage: Decimal | None
    source_url: str
    reason: str = ""  # populated when the wage is None
    warning: str = ""


@dataclass
class SourceResult:
    year: int | None
    observations: dict[str, WageObservation]
    description: str
    errors: list[str] = field(default_factory=list)


def parse_decimal(raw: Any) -> Decimal | None:
    """Parse a numeric wage cell. Returns None for BLS suppression/footnote markers."""
    if raw is None:
        return None
    text = str(raw).strip().replace(",", "").replace("$", "")
    if not text or text in {"*", "**", "#", "-", "--", "(X)", "N/A"}:
        return None
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    return value if value.is_finite() else None


# --------------------------------------------------------------------------- API


def series_id(area_code: str, soc_code: str, datatype: str) -> str:
    """OEWS series: OEUM + area(7) + industry(6, cross-industry) + SOC(6) + datatype(2)."""
    return f"OEUM{area_code.zfill(7)}000000{soc_code.replace('-', '')}{datatype}"


def _default_post(url: str, payload: dict[str, Any]) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            response = requests.post(url, json=payload, timeout=45)
            if response.status_code >= 500:
                raise SourceError(f"BLS API returned HTTP {response.status_code}")
            response.raise_for_status()
            return response.json()
        except (requests.RequestException, SourceError, ValueError) as exc:
            last_error = exc
            time.sleep(2**attempt)
    raise SourceError(f"BLS API request failed after retries: {type(last_error).__name__}")


def _values_by_year(series: dict[str, Any]) -> dict[int, Decimal | None]:
    out: dict[int, Decimal | None] = {}
    for point in series.get("data", []):
        if point.get("period") != "A01":
            continue
        try:
            year = int(point["year"])
        except (KeyError, ValueError):
            continue
        out[year] = parse_decimal(point.get("value"))
    return out


def fetch_from_api(
    socs: Iterable[str] = tuple(config.OCCUPATIONS),
    area_code: str = config.AREA_CODE,
    *,
    api_key: str | None = None,
    current_year: int,
    post: Callable[[str, dict[str, Any]], dict[str, Any]] = _default_post,
) -> SourceResult:
    """Fetch the latest published OEWS annual release for the area via the BLS API.

    The release year is the newest year published for any requested hourly-median
    series; it is not hardcoded. An occupation lacking a value for that year is
    reported as missing rather than falling back to an older year.
    """
    socs = list(socs)
    hourly = {s: series_id(area_code, s, config.DATATYPE_HOURLY_MEDIAN) for s in socs}
    annual = {s: series_id(area_code, s, config.DATATYPE_ANNUAL_MEDIAN) for s in socs}
    payload: dict[str, Any] = {
        "seriesid": [*hourly.values(), *annual.values()],
        "startyear": str(current_year - 4),
        "endyear": str(current_year),
    }
    url = config.API_URL_V1
    if api_key:
        payload["registrationkey"] = api_key
        url = config.API_URL_V2

    body = post(url, payload)
    if body.get("status") != "REQUEST_SUCCEEDED":
        # Do not echo the payload: it may contain the registration key.
        raise SourceError(f"BLS API did not succeed: {body.get('status')} {body.get('message')}")

    by_series: dict[str, dict[int, Decimal | None]] = {}
    for series in body.get("Results", {}).get("series", []):
        by_series[series.get("seriesID", "")] = _values_by_year(series)

    years = [y for s in hourly.values() for y in by_series.get(s, {})]
    if not years:
        raise SourceError("BLS API returned no OEWS data for the requested series")
    release_year = max(years)

    observations: dict[str, WageObservation] = {}
    for soc in socs:
        sid = hourly[soc]
        url_for_series = config.SERIES_URL_TEMPLATE.format(series_id=sid)
        wage = by_series.get(sid, {}).get(release_year)
        if wage is None:
            observations[soc] = WageObservation(
                soc, area_code, release_year, None, url_for_series,
                reason=f"no valid {release_year} median hourly wage published for area {area_code}",
            )
            continue
        annual_value = by_series.get(annual[soc], {}).get(release_year)
        warning = ""
        if annual_value is not None:
            expected = wage * config.ANNUAL_HOURS
            tolerance = max(annual_value * Decimal("0.01"), Decimal("20"))
            if abs(expected - annual_value) > tolerance:
                observations[soc] = WageObservation(
                    soc, area_code, release_year, None, url_for_series,
                    reason=(
                        f"hourly median {wage} x 2080 = {expected} disagrees with annual "
                        f"median {annual_value}; refusing to use ambiguous data"
                    ),
                )
                continue
        else:
            warning = "annual median unavailable, hourly/annual consistency check skipped"
        observations[soc] = WageObservation(
            soc, area_code, release_year, wage, url_for_series, warning=warning
        )
    return SourceResult(
        year=release_year,
        observations=observations,
        description=f"BLS Public Data API, OEWS May {release_year}, area {area_code}",
    )


# -------------------------------------------------------------------------- File


def read_table(path: Path) -> list[dict[str, str]]:
    """Read an official OEWS CSV or XLSX into rows keyed by upper-case header name."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return [
                {(k or "").strip().upper(): (v or "").strip() for k, v in row.items()}
                for row in csv.DictReader(handle)
            ]
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency listed in requirements
            raise SourceError("openpyxl is required to read XLSX files") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            names = [str(h).strip().upper() if h is not None else "" for h in header]
            if "OCC_CODE" in names and "H_MEDIAN" in names:
                return [
                    {n: ("" if v is None else str(v).strip()) for n, v in zip(names, row)}
                    for row in rows
                ]
        raise SourceError("No worksheet with OCC_CODE and H_MEDIAN columns was found")
    raise SourceError(f"Unsupported file type {suffix!r}; use an official OEWS .xlsx or .csv")


def _normalize_area(value: str) -> str:
    text = value.strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text.zfill(5) if text.isdigit() else text


def infer_year(path: Path) -> int | None:
    match = re.search(r"[Mm](20\d{2})", path.name)
    return int(match.group(1)) if match else None


def observations_from_rows(
    rows: list[dict[str, str]],
    year: int,
    source_url: str,
    socs: Iterable[str] = tuple(config.OCCUPATIONS),
    area_code: str = config.AREA_CODE,
) -> SourceResult:
    """Select Atlanta cross-industry rows and extract H_MEDIAN (never mean/annual)."""
    if not rows:
        raise SourceError("The OEWS file contained no data rows")
    for needed in ("AREA", "OCC_CODE", "H_MEDIAN"):
        if needed not in rows[0]:
            raise SourceError(f"The OEWS file is missing required column {needed}")

    area_rows = [r for r in rows if _normalize_area(r.get("AREA", "")) == area_code]
    if not area_rows:
        raise SourceError(f"No rows found for area code {area_code} in the OEWS file")
    title = area_rows[0].get("AREA_TITLE", "")
    if title and config.AREA_NAME_FRAGMENT.lower() not in title.lower():
        raise SourceError(f"Area {area_code} is titled {title!r}, expected Atlanta metro area")

    def cross_industry_all_ownership(row: dict[str, str]) -> bool:
        if row.get("I_GROUP") and row["I_GROUP"].strip().lower() != "cross-industry":
            return False
        if row.get("NAICS") and row["NAICS"].strip() != "000000":
            return False
        if row.get("OWN_CODE") and row["OWN_CODE"].strip() != "1235":
            return False
        return True

    observations: dict[str, WageObservation] = {}
    for soc in socs:
        matches = [
            r for r in area_rows if r.get("OCC_CODE") == soc and cross_industry_all_ownership(r)
        ]
        if len(matches) != 1:
            reason = (
                f"no cross-industry row for {soc} in area {area_code}"
                if not matches
                else f"{len(matches)} ambiguous rows for {soc} in area {area_code}"
            )
            observations[soc] = WageObservation(soc, area_code, year, None, source_url, reason=reason)
            continue
        raw = matches[0].get("H_MEDIAN", "")
        wage = parse_decimal(raw)
        if wage is None:
            observations[soc] = WageObservation(
                soc, area_code, year, None, source_url,
                reason=f"H_MEDIAN is suppressed or unavailable ({raw!r})",
            )
        else:
            observations[soc] = WageObservation(soc, area_code, year, wage, source_url)
    return SourceResult(
        year=year,
        observations=observations,
        description=f"Official OEWS file, May {year}, area {area_code}",
    )


def fetch_from_file(
    path: Path,
    *,
    year: int | None = None,
    source_url: str = config.DEFAULT_FILE_SOURCE_URL,
) -> SourceResult:
    resolved_year = year or infer_year(path)
    if resolved_year is None:
        raise SourceError("Cannot infer the OEWS year from the file name; pass --year")
    if not path.is_file():
        raise SourceError(f"OEWS file not found: {path}")
    return observations_from_rows(read_table(path), resolved_year, source_url)
